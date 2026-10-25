import sys
import io
import argparse
from collections import Counter
import docx
from docx.shared import Pt
import docx.oxml.simpletypes as st
from docx.oxml.ns import qn

# Bao dam xuat UTF-8 tren moi moi truong Windows console
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# Ho tro doc cac gia tri font size so thuc trong XML
orig_st_hps = st.ST_HpsMeasure.convert_from_xml
def safe_hps_convert(cls, xml_value):
    try:
        return orig_st_hps(xml_value)
    except ValueError:
        return Pt(float(xml_value) / 2.0)
st.ST_HpsMeasure.convert_from_xml = classmethod(safe_hps_convert)

def extract_rfonts(run):
    """Trich xuat thuoc tinh w:rFonts tu the XML rPr cua run."""
    rPr = run._r.rPr
    if rPr is None or rPr.rFonts is None:
        return {}
    elem = rPr.rFonts
    return {
        'ascii': elem.get(qn('w:ascii')),
        'hAnsi': elem.get(qn('w:hAnsi')),
        'cs': elem.get(qn('w:cs')),
        'eastAsia': elem.get(qn('w:eastAsia')),
    }

def get_effective_font_name(run, paragraph):
    """Xac dinh font name hieu luc cua mot run theo thu tu uu tien."""
    if run.font.name:
        return run.font.name
    rf = extract_rfonts(run)
    if rf.get('ascii'):
        return rf['ascii']
    if rf.get('cs'):
        return rf['cs']
    if paragraph and paragraph.style and paragraph.style.font and paragraph.style.font.name:
        return paragraph.style.font.name
    return None

def get_font_size_pt(run, paragraph):
    """Xac dinh co chu theo don vi pt tu run hoac paragraph style."""
    if run.font.size:
        return run.font.size.pt
    rPr = run._r.rPr
    if rPr is not None and rPr.find(qn('w:sz')) is not None:
        sz_val = rPr.find(qn('w:sz')).get(qn('w:val'))
        if sz_val and sz_val.isdigit():
            return int(sz_val) / 2.0
    if paragraph and paragraph.style and paragraph.style.font and paragraph.style.font.size:
        return paragraph.style.font.size.pt
    return None

def get_run_color(run, paragraph):
    """Xac dinh ma mau hex cua run."""
    if run.font.color and run.font.color.rgb:
        return str(run.font.color.rgb).upper()
    if run._r.rPr is not None:
        c = run._r.rPr.find(qn('w:color'))
        if c is not None:
            val = c.get(qn('w:val'))
            if val:
                return val.upper()
    if paragraph and paragraph.style and paragraph.style.font and paragraph.style.font.color and paragraph.style.font.color.rgb:
        return str(paragraph.style.font.color.rgb).upper()
    return "000000"

def analyze_document_fonts(doc_path):
    """Thong ke toan bo font chu, co chu va mau sac trong tai lieu."""
    doc = docx.Document(doc_path)
    font_counter = Counter()
    size_counter = Counter()
    style_counter = Counter()
    color_counter = Counter()

    for p in doc.paragraphs:
        style_counter[p.style.name] += 1
        for r in p.runs:
            f = get_effective_font_name(r, p)
            if f:
                font_counter[f] += 1
            sz = get_font_size_pt(r, p)
            if sz:
                size_counter[sz] += 1
            col = get_run_color(r, p)
            color_counter[col] += 1

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    style_counter[p.style.name] += 1
                    for r in p.runs:
                        f = get_effective_font_name(r, p)
                        if f:
                            font_counter[f] += 1
                        sz = get_font_size_pt(r, p)
                        if sz:
                            size_counter[sz] += 1
                        col = get_run_color(r, p)
                        color_counter[col] += 1

    return font_counter, size_counter, style_counter, color_counter

def search_text_and_inspect(doc_path, keyword, max_results=5):
    """Tim kiem doan van hoac o bang co chua tu khoa va in chi tiet dinh dang font, mau va bold."""
    doc = docx.Document(doc_path)
    results = []

    for idx, p in enumerate(doc.paragraphs):
        if keyword.lower() in p.text.lower():
            run_details = []
            for r_idx, r in enumerate(p.runs):
                run_details.append({
                    'index': r_idx,
                    'text': r.text,
                    'font_name': get_effective_font_name(r, p),
                    'font_size_pt': get_font_size_pt(r, p),
                    'bold': r.bold,
                    'italic': r.italic,
                    'color': get_run_color(r, p),
                    'xml_rfonts': extract_rfonts(r)
                })
            results.append({
                'type': 'paragraph',
                'index': idx,
                'style': p.style.name,
                'alignment': str(p.alignment) if p.alignment else 'None',
                'text': p.text[:120],
                'runs': run_details
            })
            if len(results) >= max_results:
                return results

    for t_idx, t in enumerate(doc.tables):
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    if keyword.lower() in p.text.lower():
                        run_details = []
                        for run_i, r in enumerate(p.runs):
                            run_details.append({
                                'index': run_i,
                                'text': r.text,
                                'font_name': get_effective_font_name(r, p),
                                'font_size_pt': get_font_size_pt(r, p),
                                'bold': r.bold,
                                'italic': r.italic,
                                'color': get_run_color(r, p),
                                'xml_rfonts': extract_rfonts(r)
                            })
                        results.append({
                            'type': 'table_cell',
                            'table_index': t_idx,
                            'row': r_idx,
                            'col': c_idx,
                            'paragraph_index': p_idx,
                            'style': p.style.name,
                            'text': p.text[:120],
                            'runs': run_details
                        })
                        if len(results) >= max_results:
                            return results

    return results

def get_context_around_paragraph(doc_path, target_idx, window=2):
    """Lay thong tin font cua cac doan van lien ke truoc va sau doan muc tieu."""
    doc = docx.Document(doc_path)
    start = max(0, target_idx - window)
    end = min(len(doc.paragraphs), target_idx + window + 1)
    
    context = []
    for idx in range(start, end):
        p = doc.paragraphs[idx]
        run_fonts = [get_effective_font_name(r, p) for r in p.runs]
        run_sizes = [get_font_size_pt(r, p) for r in p.runs]
        context.append({
            'index': idx,
            'is_target': (idx == target_idx),
            'style': p.style.name,
            'text_sample': p.text[:60],
            'fonts': run_fonts,
            'sizes': run_sizes
        })
    return context

def main():
    parser = argparse.ArgumentParser(description="Cong cu kiem tra font, mau sac va dinh dang trong tai lieu Word (.docx)")
    parser.add_argument("file", help="Duong dan tep .docx can kiem tra")
    parser.add_argument("--search", "-s", help="Tu khoa can tim kiem va kiem tra font")
    parser.add_argument("--index", "-i", type=int, help="Chi so doan van can soi thong tin chi tiet")
    parser.add_argument("--summary", action="store_true", help="In bao cao tong quan font va mau sac toan bo tai lieu")
    
    args = parser.parse_args()
    
    print(f"=== KIEM TRA DINH DANG TEP WORD: {args.file} ===")
    
    if args.summary or (not args.search and args.index is None):
        fonts, sizes, styles, colors = analyze_document_fonts(args.file)
        print("\n[1] Thong ke phan bo Font chu:")
        for f, cnt in fonts.most_common():
            print(f"  - Font: {f:<25} | So luong run: {cnt}")
            
        print("\n[2] Thong ke phan bo Co chu (pt):")
        for sz, cnt in sizes.most_common(10):
            print(f"  - Co chu: {sz:<6} pt | So luong run: {cnt}")
            
        print("\n[3] Thong ke phan bo Mau chu (Color Hex):")
        for col, cnt in colors.most_common(5):
            print(f"  - Mau: #{col:<10} | So luong run: {cnt}")

        print("\n[4] Cac Style su dung pho bien:")
        for st, cnt in styles.most_common(8):
            print(f"  - Style: {st:<25} | So luong doan: {cnt}")

    if args.search:
        print(f"\n[5] Ket qua tim kiem tu khoa: '{args.search}'")
        matches = search_text_and_inspect(args.file, args.search)
        if not matches:
            print("  Khong tim thay doan van phu hop voi tu khoa.")
        for m in matches:
            if m['type'] == 'paragraph':
                print(f"\n  * Doan van [Index {m['index']}] | Style: {m['style']} | Can le: {m['alignment']}")
            else:
                print(f"\n  * O bang [Bang {m['table_index']}, Hang {m['row']}, Cot {m['col']}] | Style: {m['style']}")
            print(f"    Van ban: \"{m['text']}\"")
            for r in m['runs']:
                print(f"      - Run #{r['index']}: \"{r['text']}\" -> Font: {r['font_name']} | Size: {r['font_size_pt']} pt | Mau: #{r['color']} | Bold: {r['bold']} | Italic: {r['italic']} | XML rFonts: {r['xml_rfonts']}")

    if args.index is not None:
        print(f"\n[6] Ngu canh doan van chi so {args.index} va cac doan lien ke:")
        ctx = get_context_around_paragraph(args.file, args.index)
        for item in ctx:
            marker = "==> [MUC TIEU]" if item['is_target'] else "    [LIEN KE]"
            print(f"  {marker} #{item['index']:<4} Style: {item['style']:<18} Fonts: {str(item['fonts']):<25} Sizes: {str(item['sizes'])} | Text: {item['text_sample']}")

if __name__ == '__main__':
    main()
